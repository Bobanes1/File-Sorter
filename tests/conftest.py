import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from file_sorter.config import Config, Course  # noqa: E402


def make_pdf(path: Path, text: str) -> Path:
    """Write a minimal one-page PDF containing `text` (no PDF-writing library needed)."""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({safe}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def make_pptx(path: Path, title: str) -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    prs.save(str(path))
    return path


def make_docx(path: Path, text: str) -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    document.save(str(path))
    return path


@pytest.fixture
def courses():
    return [
        Course("Programming Course", ["programming", "python", "prog1", "algorithm"]),
        Course("Linear Algebra", ["linear algebra", "matrix", "LA1"]),
        Course("Statistics", ["statistics", "probability", "stats"]),
    ]


@pytest.fixture
def config(tmp_path, courses):
    watch = tmp_path / "Downloads"
    watch.mkdir()
    return Config(
        watch_folder=str(watch),
        destination_root=str(tmp_path / "University"),
        courses=courses,
    )
