"""Pull readable text out of PDFs, slides and Word documents.

Only the first few pages/slides are read: course names and codes are almost
always on the title page, and it keeps sorting fast for large lecture files.
"""

from __future__ import annotations

import logging
from pathlib import Path

log = logging.getLogger(__name__)

MAX_PAGES = 5
MAX_SLIDES = 5
MAX_PARAGRAPHS = 60
MAX_CHARS = 20_000
PLAIN_TEXT_EXTENSIONS = {".txt", ".md"}


def extract_text(path: Path) -> str:
    """Return text from the start of the file, or "" if it can't be read."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            text = _pdf_text(path)
        elif ext == ".pptx":
            text = _pptx_text(path)
        elif ext == ".docx":
            text = _docx_text(path)
        elif ext in PLAIN_TEXT_EXTENSIONS:
            with open(path, encoding="utf-8", errors="ignore") as fh:
                text = fh.read(MAX_CHARS)
        else:
            return ""
    except Exception as exc:  # corrupt, encrypted or unsupported files
        log.debug("Could not read text from %s: %s", path, exc)
        return ""
    return text[:MAX_CHARS]


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return ""
    parts: list[str] = []
    meta = reader.metadata
    if meta:
        parts.extend(str(v) for v in (meta.title, meta.subject) if v)
    for page in reader.pages[:MAX_PAGES]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    props = prs.core_properties
    parts.extend(v for v in (props.title, props.subject) if v)
    for slide in list(prs.slides)[:MAX_SLIDES]:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    props = document.core_properties
    parts.extend(v for v in (props.title, props.subject) if v)
    parts.extend(p.text for p in document.paragraphs[:MAX_PARAGRAPHS])
    return "\n".join(parts)
